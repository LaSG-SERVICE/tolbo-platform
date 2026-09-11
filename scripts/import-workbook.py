from pathlib import Path
import json, sys
from openpyxl import load_workbook
src=Path(sys.argv[1]) if len(sys.argv)>1 else Path('data/Parametrage_du_moteur_Score_Pass_V1.xlsx')
out=Path(sys.argv[2]) if len(sys.argv)>2 else Path('data/scorepass-v1.json')
wb=load_workbook(src,data_only=True,read_only=True)
def rows(ws):
    all_rows=list(ws.iter_rows(values_only=True)); header_idx=next((i for i,r in enumerate(all_rows) if r and r[0] and ('Code' in r or 'Param_Key' in r or 'Field_Key' in r or 'Variable_ID' in r)),3); headers=list(all_rows[header_idx]); result=[]
    for r in all_rows[header_idx+1:]:
        if not any(v is not None for v in r): continue
        d={str(headers[i]).strip():r[i] for i in range(min(len(headers),len(r))) if headers[i] is not None and r[i] is not None}
        if d: result.append(d)
    return result
payload={'model_version':'V1.0','sheets':{ws.title:rows(ws) for ws in wb.worksheets}}
out.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(f'Paramétrage exporté vers {out}')
